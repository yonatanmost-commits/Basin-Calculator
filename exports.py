"""
exports.py — Excel, CAD (DXF), and PNG export functions.
Each function is self-contained: accepts results list + optional paths.
"""

import math
import os


# ─────────────────────────────────────────────────────────────────────────────
# Shared: raw dimension summary string (preserves original comma-sep values)
# ─────────────────────────────────────────────────────────────────────────────

def _raw_dim_string(dims: dict, stype: str) -> str:
    """Return a compact human-readable string of the original input dimensions."""
    if stype == 'frustum':
        # Show whichever dimension the user actually entered. For a top-basis
        # frustum length_bottom/width_bottom are None (derived), so reading them
        # would print "L=None"; use the top values and label them instead.
        if dims.get('dimension_basis') == 'top':
            lb = dims.get('length_top', '')
            wb = dims.get('width_top', '')
            ltag, wtag = 'L_top', 'W_top'
        else:
            lb = dims.get('length_bottom', '')
            wb = dims.get('width_bottom', '')
            ltag, wtag = 'L', 'W'
        h  = dims.get('total_height', '')
        sl = dims.get('slope_uniform') if not dims.get('uneven_slopes_enabled') \
             else f"N{dims.get('slope_north',0)} S{dims.get('slope_south',0)} " \
                  f"E{dims.get('slope_east',0)} W{dims.get('slope_west',0)}"
        return f"{ltag}={lb} {wtag}={wb} h={h} slope={sl}"
    if stype == 'rectangular':
        return (f"L={dims.get('length_internal','')} "
                f"W={dims.get('width_internal','')} "
                f"h={dims.get('total_height','')} "
                f"tw={dims.get('wall_thickness',0)}")
    if stype == 'circular':
        return (f"D={dims.get('diameter_internal','')} "
                f"h={dims.get('total_height','')} "
                f"tw={dims.get('wall_thickness',0)}")
    return ''


def _footprint_string(r: dict) -> str:
    g = r['geometry']
    if r['type'] == 'circular':
        return f"Ø{g['d_outer']:.2f} m"
    fx = g['outer_footprint_x']
    fy = g['outer_footprint_y']
    return f"{fx:.2f} × {fy:.2f} m"


# ─────────────────────────────────────────────────────────────────────────────
# 1. Excel Engine
# ─────────────────────────────────────────────────────────────────────────────

def export_excel(results: list, template_path: str, output_path: str = None) -> str:
    """
    Load existing Excel template, inject a summary table, save.

    Column layout (units carried in header parentheses):
      A  ID
      B  Name
      C  Type
      D  Raw Input Dimensions  (original string, reference only)
      E  Length / Diameter (m)   numeric
      F  Width (m)               numeric  (= E for circular)
      G  Footprint Area (m2)     Excel formula  =E*F or =PI()*(E/2)^2
      H  Water Volume (m3)       numeric
      I  Excavation Volume (m3)  numeric
      J  Concrete Volume (m3)    numeric  (blank for frustum / earth basins)
    """
    try:
        import openpyxl
        from openpyxl.utils import get_column_letter
    except ImportError:
        raise ImportError("openpyxl required: pip install openpyxl")

    if not os.path.exists(template_path):
        raise FileNotFoundError(f"Excel template not found: {template_path}")

    wb = openpyxl.load_workbook(template_path)
    ws = wb.active

    # Detect header row by 'ID' sentinel; clear everything below it each run.
    HEADER_MARKER = 'ID'
    header_row = None
    for i, row in enumerate(ws.iter_rows(min_col=1, max_col=1, values_only=True), start=1):
        if row[0] == HEADER_MARKER:
            header_row = i
            break

    if header_row is None:
        header_row = ws.max_row + 1
        ws.append([
            'ID', 'Name', 'Type', 'Raw Input Dimensions',
            'Length / Diameter (m)', 'Width (m)',
            'Footprint Area (m2)',
            'Water Volume (m3)',
            'Excavation Volume (m3)',
            'Concrete Volume (m3)',
        ])

    data_start = header_row + 1
    if ws.max_row >= data_start:
        ws.delete_rows(data_start, ws.max_row - header_row)

    for row_i, r in enumerate(results, start=data_start):
        g     = r['geometry']
        stype = r['type']

        if stype == 'circular':
            dim1 = round(g['d_outer'], 3)
            dim2 = round(g['d_outer'], 3)
            # Area formula: PI()*(D/2)^2  where D is in column E
            area_formula = f'=PI()*(E{row_i}/2)^2'
        else:
            dim1 = round(g['outer_footprint_x'], 3)
            dim2 = round(g['outer_footprint_y'], 3)
            area_formula = f'=E{row_i}*F{row_i}'

        v_concrete = g.get('v_concrete')
        ws.append([
            r['id'],
            r['name'],
            r['type'],
            _raw_dim_string(r['dimensions'], r['type']),
            dim1,
            dim2,
            area_formula,
            round(g['v_water'], 3),
            round(g['v_total_excavation'], 3),
            round(v_concrete, 3) if v_concrete is not None else None,
        ])

    data_end = data_start + len(results) - 1
    ws.append([
        'TOTAL', '', '', '', '', '',
        f'=SUM(G{data_start}:G{data_end})',
        f'=SUM(H{data_start}:H{data_end})',
        f'=SUM(I{data_start}:I{data_end})',
        f'=SUM(J{data_start}:J{data_end})',
    ])

    dest = output_path or template_path
    wb.save(dest)
    wb.close()
    return dest


# ─────────────────────────────────────────────────────────────────────────────
# 2. Process Excel — round-trip (read Basins sheet, write Basin Results sheet)
# ─────────────────────────────────────────────────────────────────────────────

BASINS_SHEET  = 'Basins'
RESULTS_SHEET = 'Basin Results'

BASINS_HEADERS = [
    'ID', 'Name', 'Type', 'Length / Diam (int.)', 'Width (int.)',
    'Height (m)', 'Water Depth (m)', 'Underdrain (m)', 'Wall Thick (m)', 'Slope', 'X (m)', 'Y (m)',
]

RESULTS_HEADERS = [
    'ID', 'Name', 'Type',
    'V_water (m3)', 'V_underdrain (m3)', 'V_process (m3)',
    'V_excav (m3)', 'V_concrete (m3)',
    'Footprint X (m)', 'Footprint Y (m)',
    'X start (m)', 'X end (m)', 'Y start (m)', 'Y end (m)',
]


def _safe_float(v, default=0.0):
    try:
        return float(v) if v is not None else default
    except (ValueError, TypeError):
        return default


def _basins_col_map(ws, header_row: int) -> dict:
    """Return {header_name: 0-based_index} from the header row of a Basins sheet."""
    cells = next(ws.iter_rows(min_row=header_row, max_row=header_row, values_only=True))
    return {str(v).strip(): i for i, v in enumerate(cells) if v is not None}


def read_basins_sheet(file_bytes: bytes) -> list:
    """Parse the Basins sheet from uploaded workbook bytes → list of structure dicts."""
    import io
    try:
        import openpyxl
    except ImportError:
        raise ImportError("openpyxl required: pip install openpyxl")

    wb = openpyxl.load_workbook(io.BytesIO(file_bytes))
    if BASINS_SHEET not in wb.sheetnames:
        raise ValueError(f"Sheet '{BASINS_SHEET}' not found in uploaded workbook.")
    ws = wb[BASINS_SHEET]

    # Locate header row by 'ID' sentinel in column A
    header_row = None
    for i, row in enumerate(ws.iter_rows(min_col=1, max_col=1, values_only=True), start=1):
        if str(row[0]).strip() == 'ID':
            header_row = i
            break
    if header_row is None:
        raise ValueError("Could not find header row (cell with 'ID') in Basins sheet.")

    # Map column names → indices so the reader is robust to column reordering
    # and to files created before the Underdrain column was added.
    cmap = _basins_col_map(ws, header_row)

    def _col(row, *names, default=None):
        for name in names:
            idx = cmap.get(name)
            if idx is not None and idx < len(row):
                v = row[idx]
                if v is not None:
                    return v
        return default

    structures = []
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        sid = str(_col(row, 'ID', default='')).strip()
        if not sid or sid.lower() == 'none':
            continue

        name    = str(_col(row, 'Name', default=sid))
        df_type = str(_col(row, 'Type', default='frustum')).strip().lower()
        # Accept both 'Width (int.)' (template header) and plain 'Width' (app table header)
        length  = str(_col(row, 'Length / Diam (int.)', default=''))
        width   = str(_col(row, 'Width (int.)', 'Width', default=''))
        height  = _safe_float(_col(row, 'Height (m)'),     1.0)
        depth   = _safe_float(_col(row, 'Water Depth (m)'), 0.5)
        ud      = _safe_float(_col(row, 'Underdrain (m)'),  0.0)
        wall_t  = _safe_float(_col(row, 'Wall Thick (m)'),  0.0)
        slope   = str(_col(row, 'Slope', default='1.0'))
        x_raw   = _col(row, 'X (m)')
        y_pos   = _safe_float(_col(row, 'Y (m)'), 0.0)

        internal_type = 'frustum' if df_type in ('frustum', 'uneven frustum') else df_type
        dims = {'total_height': height, 'water_depth': depth}

        if df_type in ('frustum', 'uneven frustum'):
            dims['length_bottom']     = length
            dims['width_bottom']      = width
            dims['underdrain_height'] = ud
            dims['wall_thickness']    = 0.0
            parts = [p.strip() for p in slope.split(';')]
            if len(parts) == 4:
                dims['uneven_slopes_enabled'] = True
                dims['slope_north']   = _safe_float(parts[0])
                dims['slope_south']   = _safe_float(parts[1])
                dims['slope_east']    = _safe_float(parts[2])
                dims['slope_west']    = _safe_float(parts[3])
                dims['slope_uniform'] = None
            else:
                dims['uneven_slopes_enabled'] = False
                dims['slope_uniform'] = _safe_float(parts[0], 1.0)
                dims['slope_north'] = dims['slope_south'] = dims['slope_east'] = dims['slope_west'] = None
        elif df_type == 'rectangular':
            dims['length_internal']   = length
            dims['width_internal']    = width
            dims['underdrain_height'] = ud
            dims['wall_thickness']    = wall_t
        else:  # circular
            dims['diameter_internal'] = length
            dims['underdrain_height'] = ud
            dims['wall_thickness']    = wall_t

        s = {'id': sid, 'name': name, 'type': internal_type, 'dimensions': dims}
        if x_raw is not None:
            s['x_pos'] = _safe_float(x_raw)
        s['y_pos'] = y_pos
        structures.append(s)

    return structures


def write_results_sheet(file_bytes: bytes, results: list) -> bytes:
    """Write computed results into Basin Results sheet; return updated workbook bytes."""
    import io
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
    except ImportError:
        raise ImportError("openpyxl required: pip install openpyxl")

    wb = openpyxl.load_workbook(io.BytesIO(file_bytes))

    if RESULTS_SHEET in wb.sheetnames:
        del wb[RESULTS_SHEET]
    ws = wb.create_sheet(RESULTS_SHEET)

    # Header row
    ws.append(RESULTS_HEADERS)
    header_fill = PatternFill('solid', fgColor='1A2E55')
    header_font = Font(bold=True, color='DEEEFF')
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center')

    # Data rows
    for r in results:
        g  = r['geometry']
        c  = r['coordinates']
        vc = g.get('v_concrete')
        ws.append([
            r['id'], r['name'], r['type'],
            round(g['v_water'],      3),
            round(g['v_underdrain'], 3),
            round(g['v_process'],    3),
            round(g['v_total_excavation'], 3),
            round(vc, 3) if vc is not None else None,
            round(g['outer_footprint_x'], 3),
            round(g['outer_footprint_y'], 3),
            round(c['x_start'], 3),
            round(c['x_end'],   3),
            round(c['y_start'], 3),
            round(c['y_end'],   3),
        ])

    # TOTAL row
    n = len(results)
    ds, de = 2, 2 + n - 1
    total_font = Font(bold=True)
    total_row = ws.append([
        'TOTAL', '', '',
        f'=SUM(D{ds}:D{de})', f'=SUM(E{ds}:E{de})', f'=SUM(F{ds}:F{de})',
        f'=SUM(G{ds}:G{de})', f'=SUM(H{ds}:H{de})',
        None, None, None, None, None, None,
    ])
    for cell in ws[ws.max_row]:
        cell.font = total_font

    # Column widths
    for col, width in zip('ABCDEFGHIJKL', [10, 28, 14, 14, 14, 14, 14, 14, 12, 12, 12, 12]):
        ws.column_dimensions[col].width = width

    # Write current X/Y positions back into the Basins sheet so re-uploading
    # the same file restores the layout.
    if BASINS_SHEET in wb.sheetnames:
        ws_b = wb[BASINS_SHEET]
        b_header_row = None
        for i, row in enumerate(ws_b.iter_rows(min_col=1, max_col=1, values_only=True), start=1):
            if str(row[0]).strip() == 'ID':
                b_header_row = i
                break
        if b_header_row is not None:
            cmap = _basins_col_map(ws_b, b_header_row)
            x_col = cmap.get('X (m)')
            y_col = cmap.get('Y (m)')
            if x_col is not None and y_col is not None:
                result_by_id = {r['id']: r for r in results}
                for row_cells in ws_b.iter_rows(min_row=b_header_row + 1):
                    cell_id = str(row_cells[0].value).strip() if row_cells[0].value is not None else ''
                    if cell_id in result_by_id:
                        c = result_by_id[cell_id]['coordinates']
                        row_cells[x_col].value = round(c['x_start'], 2)
                        row_cells[y_col].value = round(c['y_start'], 2)

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def create_process_template(output_path: str) -> str:
    """Create a fresh process_template.xlsx with a Basins input sheet."""
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
    except ImportError:
        raise ImportError("openpyxl required: pip install openpyxl")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = BASINS_SHEET

    ws.append(BASINS_HEADERS)
    header_fill = PatternFill('solid', fgColor='1A2E55')
    header_font = Font(bold=True, color='DEEEFF')
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center')

    # Example rows: ID, Name, Type, Length, Width, Height, WaterDepth, Underdrain, WallThick, Slope, X, Y
    examples = [
        ['B01', 'Aeration Tank',  'rectangular', 25.0, 12.0, 5.0, 3.8, 0.3, 0.4, '',    0.0, 0.0],
        ['B02', 'Digester',       'circular',    15.0, '',   6.0, 5.0, 0.0, 0.4, '',    0.0, 0.0],
        ['B03', 'Settling Basin', 'frustum',     20.0, 10.0, 5.0, 4.0, 0.0, 0.0, '1.0', 0.0, 0.0],
    ]
    for row in examples:
        ws.append(row)

    for col, width in zip('ABCDEFGHIJKL', [10, 24, 16, 20, 12, 10, 14, 13, 13, 16, 10, 10]):
        ws.column_dimensions[col].width = width

    # Placeholder for Basin Results sheet
    wb.create_sheet(RESULTS_SHEET)

    wb.save(output_path)
    return output_path


# ─────────────────────────────────────────────────────────────────────────────
# 3. CAD Engine (DXF)
# ─────────────────────────────────────────────────────────────────────────────

_DXF_LAYERS = {
    'outer':    ('SITE_OUTER_BOUNDARIES', 7),   # white/black
    'inner':    ('FLUID_CHANNELS',        5),   # blue
    'roads':    ('MAINTENANCE_ROADS',     1),   # red (standard CAD road colour)
}


def _add_dxf_layers(doc):
    for key, (name, color) in _DXF_LAYERS.items():
        doc.layers.add(name=name, color=color)


def _rect_dxf(msp, x0, y0, x1, y1, layer):
    pts = [(x0,y0),(x1,y0),(x1,y1),(x0,y1),(x0,y0)]
    msp.add_lwpolyline(pts, dxfattribs={'layer': layer, 'closed': True})


def _circle_dxf(msp, cx, cy, r, layer):
    msp.add_circle((cx, cy), r, dxfattribs={'layer': layer})


def export_cad(results: list, output_path: str = 'site_plan.dxf') -> str:
    """Generate a 2D DXF site plan with three layers. Returns the path written."""
    try:
        import ezdxf
    except ImportError:
        raise ImportError("ezdxf required: pip install ezdxf")

    doc = ezdxf.new(dxfversion='R2010')
    _add_dxf_layers(doc)
    msp = doc.modelspace()

    L_OUTER = _DXF_LAYERS['outer'][0]
    L_INNER = _DXF_LAYERS['inner'][0]
    L_ROADS = _DXF_LAYERS['roads'][0]

    for i, r in enumerate(results):
        g     = r['geometry']
        c     = r['coordinates']
        stype = r['type']
        x0, x1 = c['x_start'], c['x_end']
        y0, y1 = c['y_start'], c['y_end']
        cx, cy = (x0+x1)/2, (y0+y1)/2
        h = g['total_height']

        if stype == 'frustum':
            sl = g['slopes']
            _rect_dxf(msp, x0, y0, x1, y1, L_OUTER)
            _rect_dxf(msp, x0 + h*sl['W'], y0 + h*sl['S'],
                           x1 - h*sl['E'], y1 - h*sl['N'], L_INNER)

        elif stype == 'rectangular':
            _rect_dxf(msp, x0, y0, x1, y1, L_OUTER)
            li, wi = g['l_internal'], g['w_internal']
            _rect_dxf(msp, cx-li/2, cy-wi/2, cx+li/2, cy+wi/2, L_INNER)

        elif stype == 'circular':
            _circle_dxf(msp, cx, cy, g['d_outer']/2,    L_OUTER)
            _circle_dxf(msp, cx, cy, g['d_internal']/2, L_INNER)

        # Corridor gap as road rectangle
        if i < len(results) - 1:
            c_next = results[i+1]['coordinates']
            y_max  = max(y1, c_next['y_end'])
            _rect_dxf(msp, x1, 0, c_next['x_start'], y_max, L_ROADS)

        # Structure label
        msp.add_text(
            f"{r['id']} ({r['type']})",
            dxfattribs={'insert': (cx, cy), 'height': 0.5, 'layer': L_OUTER},
        )

    doc.saveas(output_path)
    return output_path


# ─────────────────────────────────────────────────────────────────────────────
# 3. PNG Image Capture
# ─────────────────────────────────────────────────────────────────────────────

def export_png(results: list, output_dir: str = '.') -> list:
    """
    Save the 2D blueprint and one 3D box view per structure as PNG files.
    Requires kaleido: pip install kaleido
    Returns list of file paths written.
    """
    from charts import make_2d_blueprint, make_3d_box

    os.makedirs(output_dir, exist_ok=True)
    paths = []

    # 2D blueprint
    fig2d = make_2d_blueprint(results)
    path2d = os.path.join(output_dir, 'blueprint_2d.png')
    fig2d.write_image(path2d, width=1600, height=900, scale=2)
    paths.append(path2d)

    # 3D per structure
    for r in results:
        fig3d = make_3d_box(r)
        fname = f"3d_{r['id'].replace('-', '_').lower()}.png"
        path3d = os.path.join(output_dir, fname)
        fig3d.write_image(path3d, width=1200, height=900, scale=2)
        paths.append(path3d)

    return paths
